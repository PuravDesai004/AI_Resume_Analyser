"""Extract document text and real PDF page references for analysis and chat."""
import io
from pathlib import Path
import docx
import pymupdf


def extract_document(file_bytes: bytes, filename: str):
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
            pages = []
            for index, page in enumerate(pdf, 1):
                text = page.get_text().strip()
                if text:
                    pages.append({"text": text, "page": index})
    elif extension == ".docx":
        document = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join(
            [p.text for p in document.paragraphs]
            + [" | ".join(cell.text for cell in row.cells)
               for table in document.tables for row in table.rows]
        ).strip()
        pages = [{"text": text, "page": None}] if text else []
    elif extension == ".txt":
        text = file_bytes.decode("utf-8-sig").strip()
        pages = [{"text": text, "page": None}] if text else []
    else:
        raise ValueError("Please use PDF, DOCX, or UTF-8 TXT files.")
    return "\n".join(p["text"] for p in pages).strip(), pages


def validate_doc_text(text: str) -> tuple[str, str]:
    """Returns (status, warning_msg). Status: 'Valid', 'Empty', 'Short', 'Excessive'."""
    word_count = len((text or "").split())
    if word_count == 0:
        return "Empty", "Document contains no text."
    elif word_count < 30:
        return "Short", f"Only {word_count} words (<30). May trigger Insufficient Context gate."
    elif word_count > 4000:
        return "Excessive", f"Document has {word_count} words (>4000). May increase latency."
    return "Valid", "Ready for analysis."
