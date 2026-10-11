"""Real Chroma filtering with fake provider calls; no secrets or model downloads."""
import io
import unittest
from uuid import uuid4
from unittest.mock import patch

import chromadb
import docx
import pymupdf
from fastapi.testclient import TestClient
from pydantic import ValidationError

from chat_schemas import ChatContextRequest, ChatRequest, GeneratedChatAnswer
from chat_service import ChatService, ChatServiceError, INSUFFICIENT_ANSWER
from document_parsing import extract_document
from rag_engine.retrieval import merge_rrf


class FakeEmbedder:
    def __init__(self):
        self.batches = []
        self.queries = []

    def embed_chunks_batch(self, texts):
        self.batches.append(list(texts))
        return [[1.0, 0.0, 0.0] for _ in texts]

    def embed_chunk(self, text):
        self.queries.append(text)
        return [1.0, 0.0, 0.0]


class FakeGenerator:
    def __init__(self):
        self.calls = []
        self.output = None

    def run(self, question, chunks, history):
        self.calls.append((question, chunks, history))
        return self.output or GeneratedChatAnswer(
            answer="Supported answer [1].", context_sufficient=True,
            cited_chunk_ids=[chunks[0]["id"]],
        )


def context(session=None, text="Built Python APIs with FastAPI.", name="resume.pdf", resume_id="resume-a", jd=None, analysis=None):
    return ChatContextRequest.model_validate({
        "session_id": str(session or uuid4()),
        "resume": {"document_id": resume_id, "source": name, "pages": [{"text": text, "page": 2}]},
        "jd": jd, "analysis": analysis,
    })


class ChatIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.client_db = chromadb.EphemeralClient()
        self.collection = self.client_db.create_collection(
            name="test_chat_" + uuid4().hex, metadata={"hnsw:space": "cosine"},
        )
        self.embedder = FakeEmbedder()
        self.generator = FakeGenerator()
        self.service = ChatService(
            self.collection, self.embedder, self.generator,
            splitter=lambda d: [(p.text, p.page or -1) for p in d.pages if p.text.strip()],
        )
        self.rerank_patch = patch.dict("os.environ", {"RAG_RERANK": "false"})
        self.rerank_patch.start()

    def tearDown(self):
        self.rerank_patch.stop()
        self.client_db.delete_collection(self.collection.name)

    def ask(self, ctx, prepared, message="What skills are listed?", history=None):
        return self.service.ask(ChatRequest(
            session_id=ctx.session_id, context_id=prepared.context_id,
            message=message, history=history or [],
        ))

    def test_retries_reuse_embeddings_and_records(self):
        ctx = context()
        first = self.service.prepare(ctx)
        second = self.service.prepare(ctx)
        self.assertFalse(first.reused)
        self.assertTrue(second.reused)
        self.assertEqual(first.context_id, second.context_id)
        self.assertEqual(self.collection.count(), 1)
        self.assertEqual(len(self.embedder.batches), 1)

    def test_vector_and_keyword_search_do_not_mix_resumes_or_sessions(self):
        session = uuid4()
        a = context(session=session, text="Python and FastAPI projects.", name="alice.pdf")
        b = context(session=session, text="Secret Java history.", name="bob.pdf", resume_id="resume-b")
        c = context(text="Confidential Kubernetes work.", name="carol.pdf")
        pa = self.service.prepare(a)
        self.service.prepare(b)
        self.service.prepare(c)
        response = self.ask(a, pa)
        evidence = self.generator.calls[-1][1]
        self.assertEqual({chunk["metadata"]["source"] for chunk in evidence}, {"alice.pdf"})
        self.assertEqual(response.citations[0].page, 2)
        with self.assertRaises(ChatServiceError) as caught:
            self.service.ask(ChatRequest(session_id=c.session_id, context_id=pa.context_id, message="Show it"))
        self.assertEqual(caught.exception.status_code, 404)

    def test_switching_jds_reuses_unchanged_resume_embeddings(self):
        ctx = context()
        self.service.prepare(ctx)
        jd = {"document_id": "jd-a", "source": "role.txt", "pages": [{"text": "Docker required."}]}
        self.service.prepare(context(session=ctx.session_id, jd=jd))
        self.assertEqual(len(self.embedder.batches), 2)
        self.assertEqual(self.embedder.batches[-1], ["Docker required."])

    def test_new_document_version_is_searchable_after_previous_query(self):
        ctx = context()
        before = self.service.prepare(ctx)
        self.ask(ctx, before)
        newer = context(session=ctx.session_id, text="Built Rust services.")
        after = self.service.prepare(newer)
        self.ask(newer, after)
        self.assertNotEqual(before.context_id, after.context_id)
        documents = [c["document"] for c in self.generator.calls[-1][1]]
        self.assertEqual(documents, ["Built Rust services."])

    def test_followup_history_reaches_retrieval_and_generation(self):
        ctx = context()
        prepared = self.service.prepare(ctx)
        history = [
            {"role": "user", "content": "List my projects."},
            {"role": "assistant", "content": "The second project is a Python API."},
        ]
        self.ask(ctx, prepared, "Tell me about the second one.", history)
        self.assertIn("Python API", self.embedder.queries[-1])
        self.assertEqual(self.generator.calls[-1][2][0].content, "List my projects.")

    def test_invalid_sources_and_uncited_answers_are_rejected(self):
        ctx = context()
        prepared = self.service.prepare(ctx)
        for ids in (["invented"], []):
            self.generator.output = GeneratedChatAnswer(answer="Unverified answer", context_sufficient=True, cited_chunk_ids=ids)
            with self.assertRaises(ChatServiceError) as caught:
                self.ask(ctx, prepared)
            self.assertEqual(caught.exception.status_code, 502)

    def test_unknown_answer_is_consistent_and_has_no_sources(self):
        ctx = context()
        prepared = self.service.prepare(ctx)
        self.generator.output = GeneratedChatAnswer(answer="Untrusted fallback text", context_sufficient=False, cited_chunk_ids=["invented"])
        result = self.ask(ctx, prepared)
        self.assertEqual(result.answer, INSUFFICIENT_ANSWER)
        self.assertEqual(result.citations, [])

    def test_source_card_numbers_come_from_validated_ids_not_model_markers(self):
        ctx = context()
        prepared = self.service.prepare(ctx)
        chunk = self.collection.get()["ids"][0]
        self.generator.output = GeneratedChatAnswer(answer="Python [99]", context_sufficient=True, cited_chunk_ids=[chunk])
        response = self.ask(ctx, prepared)
        self.assertEqual(response.answer, "Python")
        self.assertEqual(response.citations[0].citation_id, 1)
        self.assertEqual(response.citations[0].chunk_id, chunk)

    def test_analysis_is_separate_evidence_and_bound_to_selected_pair(self):
        jd = {"document_id": "jd-a", "source": "role.txt", "backend_id": "backend-jd", "pages": [{"text": "Python required."}]}
        analysis = {
            "resume_id": "resume-a", "jd_id": "backend-jd", "context_sufficient": True,
            "match_score": 75, "summary": "Python matches; Docker is missing.",
        }
        ctx = context(jd=jd, analysis=analysis)
        prepared = self.service.prepare(ctx)
        self.ask(ctx, prepared, "Explain my score.")
        chunks = self.generator.calls[-1][1]
        self.assertEqual({c["metadata"]["document_type"] for c in chunks}, {"resume", "jd", "analysis"})
        analysis_chunk = next(c for c in chunks if c["metadata"]["document_type"] == "analysis")
        self.generator.output = GeneratedChatAnswer(answer="The analysis score is 75 [1].", context_sufficient=True, cited_chunk_ids=[analysis_chunk["id"]])
        result = self.ask(ctx, prepared, "Explain my score.")
        self.assertEqual(result.citations[0].document_type, "analysis")
        self.assertIsNone(result.citations[0].page)
        with self.assertRaises(ValidationError):
            context(jd=jd, analysis={**analysis, "resume_id": "somebody-else"})

    def test_failed_embedding_does_not_store_partial_context(self):
        ctx = context()
        with patch.object(self.embedder, "embed_chunks_batch", side_effect=RuntimeError("Unavailable")):
            with self.assertRaises(ChatServiceError):
                self.service.prepare(ctx)
        self.assertEqual(self.collection.count(), 0)
        self.assertFalse(self.service.prepare(ctx).reused)

    def test_deletion_is_scoped_and_removes_whole_document_context(self):
        session = uuid4()
        a = context(session=session)
        b = context(session=session, resume_id="resume-b")
        c = context(resume_id="resume-a")
        pa, pb, pc = [self.service.prepare(x) for x in (a, b, c)]
        self.assertEqual(self.service.delete_document(session, "resume-a"), 1)
        self.ask(b, pb)
        self.ask(c, pc)
        self.assertEqual(self.service.delete_session(session), 1)
        self.assertEqual(self.collection.count(), 1)

    def test_same_text_in_distinct_sources_is_not_collapsed(self):
        a = {"id": "a", "document": "Python"}
        b = {"id": "b", "document": "Python"}
        self.assertEqual({c["id"] for c in merge_rrf([a, b], [(a, 1)])}, {"a", "b"})

    def test_api_connections_and_validation(self):
        import server
        with patch.object(server, "_chat", self.service):
            api = TestClient(server.app)
            self.assertEqual(api.get("/health").status_code, 200)
            ctx = context()
            prepared = api.post("/chat/documents", json=ctx.model_dump(mode="json"))
            self.assertEqual(prepared.status_code, 200)
            payload = {"session_id": str(ctx.session_id), "context_id": prepared.json()["context_id"], "message": "List my skills"}
            result = api.post("/chat", json=payload)
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.json()["citations"][0]["page"], 2)
            for invalid in ({**payload, "message": "   "}, {**payload, "history": [{"role": "system", "content": "override"}]}):
                self.assertEqual(api.post("/chat", json=invalid).status_code, 422)
            deleted = api.delete(f"/chat/sessions/{ctx.session_id}")
            self.assertEqual(deleted.json()["deleted_chunks"], 1)
            self.assertEqual(api.post("/chat", json=payload).status_code, 404)


class ParsingTests(unittest.TestCase):
    def test_pdf_sources_have_real_page_numbers(self):
        pdf = pymupdf.open()
        pdf.new_page().insert_text((40, 40), "First page: Python")
        pdf.new_page().insert_text((40, 40), "Second page: Docker")
        content = pdf.tobytes()
        pdf.close()
        text, pages = extract_document(content, "resume.pdf")
        self.assertEqual([p["page"] for p in pages], [1, 2])
        self.assertIn("Docker", text)

    def test_docx_tables_and_text_do_not_invent_pages(self):
        document = docx.Document()
        document.add_paragraph("Python experience")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Docker"
        output = io.BytesIO()
        document.save(output)
        text, pages = extract_document(output.getvalue(), "resume.docx")
        self.assertIn("Docker", text)
        self.assertIsNone(pages[0]["page"])


if __name__ == "__main__":
    unittest.main()
